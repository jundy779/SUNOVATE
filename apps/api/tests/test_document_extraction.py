"""Bounded extraction tests for explicitly simulated files."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Sequence

import pytest
import xlwt
from docx import Document
from openpyxl import Workbook

from sunovate_api.hermes.chunking import chunk_text
from sunovate_api.hermes.document_extraction import (
    MAX_FILE_BYTES,
    ExtractionRejected,
    extract_simulated_document,
)


def _make_text_pdf(pages: Sequence[str]) -> bytes:
    """Build a tiny deterministic text PDF without adding a test dependency."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    page_ids = " ".join(f"{index + 3} 0 R" for index in range(len(pages)))
    objects.append(
        f"<< /Type /Pages /Kids [{page_ids}] /Count {len(pages)} >>".encode()
    )
    font_id = 3 + len(pages)
    content_start = font_id + 1
    for index in range(len(pages)):
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_start + index} 0 R >>"
            ).encode()
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page in pages:
        safe_text = page.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode()
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream"
        )

    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, value in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{object_id} 0 obj\n".encode() + value + b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(payload)


def test_extracts_utf8_text_without_persisting_source() -> None:
    payload = b"Pasal 1\nIsi simulasi yang cukup panjang untuk diperiksa oleh sistem."

    result = extract_simulated_document(
        filename="SIMULASI-contoh.txt", content_type="text/plain", payload=payload
    )

    assert result.text.startswith("Pasal 1")
    assert result.extension == ".txt"
    assert result.size_bytes == len(payload)
    assert len(result.checksum_sha256) == 64


def test_extracts_representative_pdf_pages() -> None:
    payload = _make_text_pdf(
        [
            "SIMULASI halaman satu berisi ringkasan pemeriksaan yang aman.",
            "SIMULASI halaman dua wajib ditinjau manusia sebelum diterbitkan.",
        ]
    )

    result = extract_simulated_document(
        filename="SIMULASI-laporan.pdf", content_type="application/pdf", payload=payload
    )

    assert result.page_count == 2
    assert "halaman satu" in result.text
    assert "halaman dua" in result.text
    assert result.checksum_sha256


def test_extracts_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("Pendahuluan", level=1)
    document.add_paragraph("SIMULASI dokumen untuk pemeriksaan isi yang aman dan terbatas.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Kolom satu"
    table.cell(0, 1).text = "Kolom dua"
    document.add_paragraph("Kesimpulan setelah tabel wajib tetap berada pada urutan aslinya.")
    buffer = io.BytesIO()
    document.save(buffer)

    result = extract_simulated_document(
        filename="SIMULASI-dokumen.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        payload=buffer.getvalue(),
    )

    assert result.text.startswith("# Pendahuluan")
    assert result.text.index("Kolom satu | Kolom dua") < result.text.index("Kesimpulan")
    assert result.paragraph_count == 4
    chunks = chunk_text(result.text, "REPORT", document_label="SIMULASI laporan")
    assert chunks[0].section_label == "Pendahuluan"


def test_extracts_xlsx_values_without_executing_formulas() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Hasil Lelang"
    worksheet.append(["Seri", "Nominal", "Catatan"])
    worksheet.append(["SIMULASI-FR", 1_000_000, "Data untuk pemeriksaan manusia"])
    worksheet.append(["Rumus", "=1+1", None])
    hidden = workbook.create_sheet("Catatan Tersembunyi")
    hidden.sheet_state = "hidden"
    hidden.append(["SIMULASI catatan tersembunyi tetap dibaca untuk pemeriksaan keamanan"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = extract_simulated_document(
        filename="SIMULASI-hasil-lelang.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        payload=buffer.getvalue(),
    )

    assert "Lembar: Hasil Lelang" in result.text
    assert "SIMULASI-FR | 1000000 | Data untuk pemeriksaan manusia" in result.text
    assert "=1+1" not in result.text
    assert "Lembar: Catatan Tersembunyi" in result.text
    assert "catatan tersembunyi tetap dibaca" in result.text
    assert result.sheet_count == 2


def test_extracts_legacy_xls_values() -> None:
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Rekonsiliasi")
    rows = [
        ("Seri", "Status"),
        ("SIMULASI-FR", "Perlu pemeriksaan manusia"),
    ]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            worksheet.write(row_index, column_index, value)
    notes = workbook.add_sheet("Catatan")
    notes.write(0, 0, "SIMULASI lembar kedua untuk pemeriksaan struktur")
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = extract_simulated_document(
        filename="SIMULASI-rekonsiliasi.xls",
        content_type="application/vnd.ms-excel",
        payload=buffer.getvalue(),
    )

    assert "Lembar: Rekonsiliasi" in result.text
    assert "SIMULASI-FR | Perlu pemeriksaan manusia" in result.text
    assert "Lembar: Catatan" in result.text
    assert result.sheet_count == 2


@pytest.mark.parametrize(
    ("filename", "content_type", "payload", "message"),
    [
        ("dokumen.txt", "text/plain", b"Isi simulasi yang panjang dan aman.", "diawali"),
        ("SIMULASI.exe", "application/octet-stream", b"data", "XLSX, atau XLS"),
        ("SIMULASI.txt", "text/plain", b"teks\x00biner yang tidak aman", "data biner"),
        ("SIMULASI.pdf", "application/pdf", b"bukan pdf yang valid", "format PDF"),
        ("SIMULASI.xls", "application/vnd.ms-excel", b"bukan xls", "format XLS"),
        ("SIMULASI.txt", "text/plain", b"a" * (MAX_FILE_BYTES + 1), "maksimal 2 MB"),
    ],
)
def test_rejects_unsafe_or_mislabelled_files(
    filename: str, content_type: str, payload: bytes, message: str
) -> None:
    with pytest.raises(ExtractionRejected, match=message):
        extract_simulated_document(
            filename=filename, content_type=content_type, payload=payload
        )


def test_rejects_macro_payload_before_docx_parsing() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", "<document />")
        archive.writestr("word/vbaProject.bin", b"macro")

    with pytest.raises(ExtractionRejected, match="makro") as rejected:
        extract_simulated_document(
            filename="SIMULASI-makro.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            payload=buffer.getvalue(),
        )
    assert rejected.value.extraction_status == "QUARANTINED"
    assert rejected.value.reason_code == "MACRO_DETECTED"
