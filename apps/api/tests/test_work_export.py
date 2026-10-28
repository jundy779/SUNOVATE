"""Excel export of work-monitoring data (TASK-033). SIMULATED_DATA only."""

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from sunovate_api.modules.exports import build_workbook
from sunovate_api.testing import login, make_client

API = "/api/v1/exports/work"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(content: bytes):  # type: ignore[no-untyped-def]
    return load_workbook(io.BytesIO(content)).active


def _titles(sheet) -> list[str]:  # type: ignore[no-untyped-def]
    return [row[0] for row in sheet.iter_rows(min_row=6, values_only=True) if row[0]]


@pytest.mark.parametrize("kind", ["outputs", "tasks", "milestones", "pending", "followUps"])
def test_export_returns_filtered_workbook_with_autofilter(kind: str) -> None:
    client = make_client()
    response = client.get(API, params={"kind": kind}, headers=login(client, "sim.staff.pppsun"))
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"].startswith('attachment; filename="SUNOVATE-')
    assert "no-store" in response.headers["cache-control"]
    sheet = _sheet(response.content)
    assert [cell.value for cell in sheet[5]][:3] == ["Judul", "Tim", "Status"]
    assert sheet.auto_filter.ref is not None and sheet.auto_filter.ref.startswith("A5:")
    assert sheet.freeze_panes == "A6"
    notice = str(sheet["A3"].value)
    assert notice.startswith("Data contoh.")
    assert "SIMULATED_DATA" not in notice  # user-facing text has no internal labels


def test_export_contains_exactly_the_rows_the_user_can_read() -> None:
    client = make_client()
    for user in ("sim.staff.pppsun", "sim.coordinator"):
        headers = login(client, user)
        listed = {
            item["title"]
            for item in client.get("/api/v1/tasks?pageSize=100", headers=headers).json()["data"]
        }
        response = client.get(API, params={"kind": "tasks"}, headers=headers)
        exported = set(_titles(_sheet(response.content)))
        assert exported == listed, user


def test_export_is_audited_without_row_content() -> None:
    client = make_client()
    client.get(API, params={"kind": "tasks"}, headers=login(client, "sim.staff.pppsun"))
    logs = client.get(
        "/api/v1/audit-logs?action=WORK_EXPORT", headers=login(client, "sim.admin")
    ).json()["data"]
    assert logs and logs[0]["detail"]["kind"] == "tasks"
    assert set(logs[0]["detail"]) == {"kind", "rows"}


def test_export_requires_authentication_permission_and_known_kind() -> None:
    client = make_client()
    assert client.get(API, params={"kind": "tasks"}).status_code == 401
    staff = login(client, "sim.staff.pppsun")
    assert client.get(API, params={"kind": "secure-documents"}, headers=staff).status_code == 422
    assert client.get(API, headers=staff).status_code == 422


def test_user_text_never_becomes_a_formula() -> None:
    content = build_workbook(
        title="Tugas",
        rows=[["=HYPERLINK(\"http://contoh\")", "PPSUN", "+1", None, "", None, "-", "@SUM(A1)"]],
        exported_by="Uji",
        exported_at=datetime(2026, 10, 22, 9, 0),
    )
    sheet = _sheet(content)
    for cell in sheet[6]:
        if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
            assert cell.data_type == "s", cell.value



def test_attention_export_contains_every_urgent_item_in_scope() -> None:
    client = make_client()
    for user in ("sim.staff.pppsun", "sim.coordinator"):
        headers = login(client, user)
        summary = client.get("/api/v1/dashboard/summary", headers=headers).json()["data"]
        response = client.get("/api/v1/exports/attention", headers=headers)
        assert response.status_code == 200
        assert "SUNOVATE-perlu-ditindaklanjuti-" in response.headers["content-disposition"]
        sheet = _sheet(response.content)
        assert [cell.value for cell in sheet[5]][:2] == ["Uraian", "Jenis"]
        titles = _titles(sheet)
        assert len(titles) == summary["metrics"]["attentionCount"], user
        assert {item["title"] for item in summary["attention"]} <= set(titles)


def test_attention_export_requires_authentication() -> None:
    assert make_client().get("/api/v1/exports/attention").status_code == 401
