"""Evaluate bounded document extraction with an in-memory SIMULATED_DATA corpus.

The command reports aggregate quality signals and case identifiers only. Source text,
filenames, checksums, and extracted content are deliberately omitted from its output.
"""

from __future__ import annotations

import io
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import xlwt
from docx import Document
from openpyxl import Workbook

from ..hermes.document_extraction import (
    ExtractedDocument,
    ExtractionRejected,
    extract_simulated_document,
)
from ..hermes.scope_guard import is_instruction_like

FIXTURE = (
    Path(__file__).resolve().parents[4]
    / "tests"
    / "fixtures"
    / "document_extraction_set.json"
)


@dataclass(frozen=True)
class CaseResult:
    case_id: str
    extracted: bool
    fragment_hits: int
    fragment_total: int
    forbidden_hits: int
    structure_passed: bool
    instruction_classification_passed: bool


@dataclass(frozen=True)
class ExtractionMetrics:
    cases: int
    extraction_success_rate: float
    fragment_recall: float
    forbidden_hits: int
    structure_pass_rate: float
    instruction_classification_accuracy: float
    passed: bool
    case_results: tuple[CaseResult, ...]


def load_evaluation_set() -> dict[str, Any]:
    payload: dict[str, Any] = json.loads(FIXTURE.read_text(encoding="utf-8"))
    if payload.get("dataLabel") != "SIMULATED_DATA":
        raise ValueError("Dataset evaluasi ekstraksi wajib berlabel SIMULATED_DATA")
    cases = payload.get("cases")
    if not isinstance(cases, list) or not cases:
        raise ValueError("Dataset evaluasi ekstraksi wajib memiliki kasus")
    return payload


def _text_pdf(pages: list[str]) -> bytes:
    """Build a minimal deterministic PDF without another runtime dependency."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    page_ids = " ".join(f"{index + 3} 0 R" for index in range(len(pages)))
    objects.append(
        f"<< /Type /Pages /Kids [{page_ids}] /Count {len(pages)} >>".encode()
    )
    font_id = 3 + len(pages)
    content_start = font_id + 1
    for index in range(len(pages)):
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_start + index} 0 R >>"
            ).encode()
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page in pages:
        safe_text = page.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode()
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream"
        )

    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, value in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{object_id} 0 obj\n".encode() + value + b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(payload)


def _docx(source: dict[str, Any]) -> bytes:
    document = Document()
    for block in source["blocks"]:
        if block["type"] == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif block["type"] == "paragraph":
            document.add_paragraph(block["text"])
        elif block["type"] == "table":
            rows = block["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = str(value)
        else:
            raise ValueError(f"Blok DOCX tidak dikenal: {block['type']}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx(source: dict[str, Any]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_data in source["sheets"]:
        worksheet = workbook.create_sheet(sheet_data["name"])
        for row in sheet_data["rows"]:
            worksheet.append(row)
        if sheet_data.get("hidden", False):
            worksheet.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xls(source: dict[str, Any]) -> bytes:
    workbook = xlwt.Workbook()
    for sheet_data in source["sheets"]:
        worksheet = workbook.add_sheet(sheet_data["name"])
        for row_index, row in enumerate(sheet_data["rows"]):
            for column_index, value in enumerate(row):
                worksheet.write(row_index, column_index, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_payload(case: dict[str, Any]) -> bytes:
    source: dict[str, Any] = case["source"]
    file_format = case["format"]
    if file_format in {"txt", "md"}:
        text = source.get("text")
        if not isinstance(text, str):
            raise ValueError("Sumber teks evaluasi wajib berupa string")
        return text.encode("utf-8")
    if file_format == "pdf":
        return _text_pdf(source["pages"])
    if file_format == "docx":
        return _docx(source)
    if file_format == "xlsx":
        return _xlsx(source)
    if file_format == "xls":
        return _xls(source)
    raise ValueError(f"Format evaluasi tidak didukung: {file_format}")


def _structure_passed(case: dict[str, Any], result: ExtractedDocument) -> bool:
    expected = case.get("expectedStructure", {})
    checks: list[bool] = []
    if "pageCount" in expected:
        checks.append(result.page_count == expected["pageCount"])
    if "paragraphCount" in expected:
        checks.append(result.paragraph_count == expected["paragraphCount"])
    if "sheetCount" in expected:
        checks.append(result.sheet_count == expected["sheetCount"])
    position = -1
    for fragment in case.get("expectedSequence", []):
        next_position = result.text.find(fragment, position + 1)
        checks.append(next_position > position)
        position = next_position
    return all(checks)


def evaluate() -> ExtractionMetrics:
    payload = load_evaluation_set()
    results: list[CaseResult] = []
    total_fragments = 0
    fragment_hits = 0
    forbidden_hits = 0

    for case in payload["cases"]:
        required = case.get("expectedFragments", [])
        forbidden = case.get("forbiddenFragments", [])
        total_fragments += len(required)
        try:
            result = extract_simulated_document(
                filename=case["filename"],
                content_type=case["contentType"],
                payload=build_payload(case),
            )
        except ExtractionRejected:
            results.append(
                CaseResult(
                    case_id=case["id"],
                    extracted=False,
                    fragment_hits=0,
                    fragment_total=len(required),
                    forbidden_hits=0,
                    structure_passed=False,
                    instruction_classification_passed=False,
                )
            )
            continue

        folded_text = result.text.casefold()
        case_fragment_hits = sum(fragment.casefold() in folded_text for fragment in required)
        case_forbidden_hits = sum(fragment.casefold() in folded_text for fragment in forbidden)
        fragment_hits += case_fragment_hits
        forbidden_hits += case_forbidden_hits
        instruction_passed = is_instruction_like(result.text) == case.get(
            "expectedInstructionLike", False
        )
        results.append(
            CaseResult(
                case_id=case["id"],
                extracted=True,
                fragment_hits=case_fragment_hits,
                fragment_total=len(required),
                forbidden_hits=case_forbidden_hits,
                structure_passed=_structure_passed(case, result),
                instruction_classification_passed=instruction_passed,
            )
        )

    cases = len(results)
    extraction_rate = sum(result.extracted for result in results) / cases
    recall = fragment_hits / total_fragments if total_fragments else 0.0
    structure_rate = sum(result.structure_passed for result in results) / cases
    instruction_accuracy = (
        sum(result.instruction_classification_passed for result in results) / cases
    )
    thresholds = payload["thresholds"]
    passed = (
        extraction_rate >= thresholds["extractionSuccessRate"]
        and recall >= thresholds["fragmentRecall"]
        and forbidden_hits == thresholds["forbiddenHits"]
        and structure_rate >= thresholds["structurePassRate"]
        and instruction_accuracy >= thresholds["instructionClassificationAccuracy"]
    )
    return ExtractionMetrics(
        cases=cases,
        extraction_success_rate=round(extraction_rate, 4),
        fragment_recall=round(recall, 4),
        forbidden_hits=forbidden_hits,
        structure_pass_rate=round(structure_rate, 4),
        instruction_classification_accuracy=round(instruction_accuracy, 4),
        passed=passed,
        case_results=tuple(results),
    )


def main() -> int:
    metrics = evaluate()
    print(json.dumps(asdict(metrics), indent=2, ensure_ascii=False))
    return 0 if metrics.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
