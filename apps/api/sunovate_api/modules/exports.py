"""Excel download of work-monitoring data (TASK-033).

Rows are exactly what the caller may already read in the application (same authorization and
classification filter as the list endpoints). Every download is audited. RAHASIA documents are
never part of these exports: they live in the secure vault, which has no export route.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Annotated, Any, Final, Literal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..authz.dependencies import Ctx, require
from ..authz.principal import Principal
from ..context import AppContext
from ..domain import deadline_state
from ..secure.vault import jakarta_now
from ..store.records import FollowUp, Governed, Milestone, PendingMatter, Task, WorkOutput
from .dashboard import attention_items

router = APIRouter(prefix="/exports", tags=["exports"])

Kind = Literal["outputs", "tasks", "milestones", "pending", "followUps"]
MAX_ROWS: Final = 5000
XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_KINDS: Final[dict[str, tuple[type[Governed], str, str, str]]] = {
    # kind: (record type, permission, sheet title, file slug)
    "outputs": (WorkOutput, "work:read", "Output kerja", "output-kerja"),
    "tasks": (Task, "work:read", "Tugas", "tugas"),
    "milestones": (Milestone, "work:read", "Tahapan penting", "tahapan-penting"),
    "pending": (PendingMatter, "pending:read", "Perlu tindak lanjut", "perlu-tindak-lanjut"),
    "followUps": (FollowUp, "pending:read", "Tindak lanjut", "tindak-lanjut"),
}
_STATUS: Final = {
    "NOT_STARTED": "Belum dimulai",
    "IN_PROGRESS": "Berjalan",
    "WAITING_INTERNAL": "Menunggu internal",
    "WAITING_EXTERNAL": "Menunggu eksternal",
    "ON_HOLD": "Ditunda",
    "COMPLETED": "Selesai",
    "CANCELLED": "Dibatalkan",
    "OPEN": "Terbuka",
    "RESOLVED": "Selesai",
    "DONE": "Selesai",
    "PLANNED": "Direncanakan",
    "ACHIEVED": "Tercapai",
    "MISSED": "Tidak tercapai",
}
_DEADLINE: Final = {
    "SAFE": "Aman",
    "DUE_SOON": "Jatuh tempo dekat",
    "DUE_TODAY": "Jatuh tempo hari ini",
    "OVERDUE": "Terlambat",
    "COMPLETED": "Selesai",
}
_PRIORITY: Final = {"LOW": "Rendah", "NORMAL": "Normal", "HIGH": "Tinggi", "URGENT": "Mendesak"}
_HEADERS: Final = (
    "Judul",
    "Tim",
    "Status",
    "Progres (%)",
    "Prioritas",
    "Tanggal target/tenggat",
    "Status tenggat",
    "Klasifikasi",
)
_NAVY = "0D2E5E"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _due(record: Governed) -> date | None:
    value = getattr(record, "target_date", None) if isinstance(record, WorkOutput) else None
    return value if value is not None else getattr(record, "due_date", None)


def _row(ctx: AppContext, record: Governed, labels: dict[str, str]) -> list[Any]:
    due = _due(record)
    status = str(getattr(record, "status", ""))
    state = deadline_state(due, status, ctx.clock(), ctx.settings.deadline_due_soon_days)
    progress = getattr(record, "progress", None)
    priority = getattr(record, "priority", None)
    return [
        str(getattr(record, "title", "")),
        ctx.store.team_code(record.owner_team_id) or "-",
        _STATUS.get(status, status.replace("_", " ").capitalize()),
        progress if isinstance(progress, int) else None,
        _PRIORITY.get(str(priority), "") if priority else "",
        due,
        _DEADLINE.get(state or "", "-"),
        labels.get(record.classification_code, record.classification_code),
    ]


def _as_date(value: Any) -> date | None:
    return date.fromisoformat(value) if isinstance(value, str) else value


def _as_text(cell: Cell) -> None:
    """Never let user-entered text become an Excel formula (formula/CSV injection)."""
    if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
        cell.data_type = "s"


def build_workbook(
    *,
    title: str,
    rows: list[list[Any]],
    exported_by: str,
    exported_at: datetime,
    headers: tuple[str, ...] = _HEADERS,
    widths: tuple[int, ...] = (48, 10, 20, 12, 12, 20, 20, 14),
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title[:31]
    sheet["A1"] = f"SUNOVATE - {title}"
    sheet["A1"].font = Font(bold=True, size=14, color=_NAVY)
    sheet["A2"] = f"Diunduh oleh {exported_by} pada {exported_at:%d-%m-%Y %H:%M} WIB"
    sheet["A3"] = (
        "Data contoh (SIMULATED_DATA). Isi hanya mencakup data yang boleh dilihat pengunduh; "
        "jangan diteruskan kepada pihak yang tidak berwenang."
    )
    for cell in (sheet["A2"], sheet["A3"]):
        cell.font = Font(italic=True, color="6B7280")

    header_row = 5
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=header_row + offset, column=column, value=value)
            _as_text(cell)
            if isinstance(value, date):
                cell.number_format = "DD-MM-YYYY"

    last_row = header_row + max(len(rows), 1)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    sheet.freeze_panes = f"A{header_row + 1}"
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/work", response_class=Response)
def export_work(
    ctx: Ctx,
    principal: Annotated[Principal, Depends(require("work:read", resource="work_export"))],
    kind: Annotated[Kind, Query()],
) -> Response:
    record_type, permission, title, slug = _KINDS[kind]
    ctx.authz.require(principal, permission, resource="work_export")
    records = ctx.authz.list_readable(principal, record_type)
    records.sort(key=lambda r: (_due(r) is None, _due(r) or date.max))
    records = records[:MAX_ROWS]
    labels = {c.code: c.label for c in ctx.store.classifications()}
    rows = [_row(ctx, record, labels) for record in records]
    now = jakarta_now()
    content = build_workbook(
        title=title, rows=rows, exported_by=principal.display_name, exported_at=now
    )
    filename = f"SUNOVATE-{slug}-{now:%Y%m%d}.xlsx"
    return _download(ctx, principal, kind, len(rows), content, filename)


def _download(
    ctx: AppContext, principal: Principal, kind: str, rows: int, content: bytes, filename: str
) -> Response:
    with ctx.store.transaction():
        ctx.store.audit(
            actor_user_id=principal.user_id,
            action="WORK_EXPORT",
            entity_type="work_export",
            detail={"kind": kind, "rows": rows},
        )
    return Response(
        content=content,
        media_type=XLSX,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store, private, max-age=0",
            "X-Content-Type-Options": "nosniff",
        },
    )


_ATTENTION_HEADERS: Final = (
    "Uraian",
    "Jenis",
    "Tim",
    "Penanggung jawab",
    "Batas waktu",
    "Status tenggat",
    "Status pekerjaan",
)


@router.get("/attention", response_class=Response)
def export_attention(
    ctx: Ctx,
    principal: Annotated[Principal, Depends(require("work:read", resource="work_export"))],
) -> Response:
    """The dashboard's "Perlu ditindaklanjuti" list in full (the screen shows the first six)."""
    items = attention_items(ctx, principal)[:MAX_ROWS]
    rows = [
        [
            item["title"],
            item["category"],
            item["ownerTeamCode"] or "-",
            item["assigneeName"] or "Belum ditetapkan",
            _as_date(item["dueDate"]),
            _DEADLINE.get(item["deadlineState"], "-"),
            _STATUS.get(item["workStatus"], str(item["workStatus"]).replace("_", " ").capitalize()),
        ]
        for item in items
    ]
    now = jakarta_now()
    content = build_workbook(
        title="Perlu ditindaklanjuti",
        rows=rows,
        exported_by=principal.display_name,
        exported_at=now,
        headers=_ATTENTION_HEADERS,
        widths=(48, 14, 10, 28, 14, 20, 20),
    )
    filename = f"SUNOVATE-perlu-ditindaklanjuti-{now:%Y%m%d}.xlsx"
    return _download(ctx, principal, "attention", len(rows), content, filename)
