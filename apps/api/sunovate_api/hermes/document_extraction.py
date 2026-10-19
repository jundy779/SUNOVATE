"""Bounded in-memory extraction for explicitly simulated development files.

No source file is persisted, uploaded to a provider, indexed, or sent to an
external service. Production ingestion requires a separate approved design.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from hashlib import sha256
from pathlib import PurePath
from typing import Any

import xlrd
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from olefile import OleFileIO
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 100_000
MAX_PDF_PAGES = 50
MAX_DOCX_ENTRIES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_SPREADSHEET_SHEETS = 20
MAX_SPREADSHEET_ROWS = 5_000
MAX_SPREADSHEET_COLUMNS = 100
MAX_SPREADSHEET_CELLS = 50_000
MAX_XLSX_ENTRIES = 1_000
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024

_SIMULATED_NAME = re.compile(r"^simulasi(?:[-_. ]|$)", re.IGNORECASE)
_ALLOWED_TYPES: dict[str, frozenset[str]] = {
    ".txt": frozenset({"text/plain", "application/octet-stream", ""}),
    ".md": frozenset({"text/markdown", "text/plain", "application/octet-stream", ""}),
    ".pdf": frozenset({"application/pdf", "application/octet-stream", ""}),
    ".docx": frozenset(
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/octet-stream",
            "",
        }
    ),
    ".xlsx": frozenset(
        {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
            "",
        }
    ),
    ".xls": frozenset(
        {
            "application/vnd.ms-excel",
            "application/octet-stream",
            "",
        }
    ),
}


class ExtractionRejected(ValueError):
    """Expected validation failure with a safe workflow classification."""

    def __init__(
        self,
        message: str,
        *,
        extraction_status: str = "FAILED",
        reason_code: str = "EXTRACTION_FAILED",
    ) -> None:
        super().__init__(message)
        self.extraction_status = extraction_status
        self.reason_code = reason_code


def _quarantined(message: str, reason_code: str) -> ExtractionRejected:
    return ExtractionRejected(
        message,
        extraction_status="QUARANTINED",
        reason_code=reason_code,
    )


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    text: str
    extension: str
    mime_type: str
    size_bytes: int
    checksum_sha256: str
    page_count: int | None
    paragraph_count: int | None
    sheet_count: int | None


def _clean(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = "\n".join(line.rstrip() for line in normalized.split("\n"))
    return re.sub(r"\n{4,}", "\n\n\n", normalized).strip()


def _validate_text(text: str) -> str:
    cleaned = _clean(text)
    if len(cleaned) < 20:
        raise ExtractionRejected("Isi berkas terlalu singkat untuk diperiksa")
    if len(cleaned) > MAX_EXTRACTED_CHARACTERS:
        raise _quarantined(
            "Isi hasil pembacaan melebihi batas 100.000 karakter", "EXTRACTED_TEXT_LIMIT"
        )
    return cleaned


def _extract_plain(payload: bytes) -> tuple[str, None, None]:
    if b"\x00" in payload:
        raise _quarantined(
            "Berkas teks mengandung data biner yang tidak didukung", "BINARY_CONTENT"
        )
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionRejected("Berkas teks harus menggunakan penyandian UTF-8") from exc
    return text, None, None


def _extract_pdf(payload: bytes) -> tuple[str, int, None]:
    if not payload.startswith(b"%PDF-"):
        raise ExtractionRejected("Isi berkas tidak sesuai dengan format PDF")
    try:
        reader = PdfReader(io.BytesIO(payload), strict=True)
        if reader.is_encrypted:
            raise ExtractionRejected("PDF terenkripsi belum dapat diperiksa")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise _quarantined("PDF simulasi dibatasi maksimal 50 halaman", "PAGE_LIMIT")
        parts: list[str] = []
        total = 0
        for page in reader.pages:
            content = page.extract_text() or ""
            total += len(content)
            if total > MAX_EXTRACTED_CHARACTERS:
                raise _quarantined(
                    "Isi hasil pembacaan melebihi batas 100.000 karakter",
                    "EXTRACTED_TEXT_LIMIT",
                )
            parts.append(content)
        return "\n\n".join(parts), len(reader.pages), None
    except ExtractionRejected:
        raise
    except Exception as exc:
        raise ExtractionRejected("PDF tidak dapat dibaca dengan aman") from exc


def _docx_preflight(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            names = {entry.filename.lower() for entry in entries}
            if len(entries) > MAX_DOCX_ENTRIES:
                raise _quarantined(
                    "Struktur DOCX terlalu kompleks untuk tahap simulasi", "STRUCTURE_LIMIT"
                )
            if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise _quarantined(
                    "Ukuran isi DOCX setelah dibuka melebihi batas aman", "EXPANDED_SIZE_LIMIT"
                )
            if "[content_types].xml" not in names or "word/document.xml" not in names:
                raise ExtractionRejected("Isi berkas tidak sesuai dengan format DOCX")
            if any(name.endswith("vbaproject.bin") for name in names):
                raise _quarantined(
                    "Dokumen yang mengandung makro tidak diizinkan", "MACRO_DETECTED"
                )
    except ExtractionRejected:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ExtractionRejected("DOCX tidak dapat dibaca dengan aman") from exc


def _extract_docx(payload: bytes) -> tuple[str, None, int]:
    _docx_preflight(payload)
    try:
        document = Document(io.BytesIO(payload))
        blocks: list[str] = []
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue
                style_name = item.style.name if item.style is not None else ""
                heading = re.fullmatch(r"Heading\s+([1-6])", style_name, re.IGNORECASE)
                blocks.append(f"{'#' * int(heading.group(1))} {text}" if heading else text)
                continue
            if isinstance(item, Table):
                for row in item.rows:
                    value = " | ".join(cell.text.strip() for cell in row.cells)
                    if value.strip(" |"):
                        blocks.append(value)
        return "\n\n".join(blocks), None, len(blocks)
    except Exception as exc:
        raise ExtractionRejected("DOCX tidak dapat dibaca dengan aman") from exc


def _spreadsheet_text(sheets: list[tuple[str, list[list[Any]]]]) -> tuple[str, None, None, int]:
    if len(sheets) > MAX_SPREADSHEET_SHEETS:
        raise _quarantined("Berkas Excel dibatasi maksimal 20 lembar", "SHEET_LIMIT")
    parts: list[str] = []
    total_cells = 0
    for sheet_name, rows in sheets:
        parts.append(f"Lembar: {sheet_name}")
        for row_number, row in enumerate(rows, start=1):
            if row_number > MAX_SPREADSHEET_ROWS:
                raise _quarantined(
                    "Setiap lembar Excel dibatasi maksimal 5.000 baris", "ROW_LIMIT"
                )
            values = [
                str(value).strip()
                for value in row[:MAX_SPREADSHEET_COLUMNS]
                if value not in (None, "")
            ]
            total_cells += len(values)
            if total_cells > MAX_SPREADSHEET_CELLS:
                raise _quarantined(
                    "Berkas Excel dibatasi maksimal 50.000 sel berisi", "CELL_LIMIT"
                )
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts), None, None, len(sheets)


def _extract_xlsx(payload: bytes) -> tuple[str, None, None, int]:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            names = {entry.filename.lower() for entry in entries}
            if len(entries) > MAX_XLSX_ENTRIES:
                raise _quarantined(
                    "Struktur XLSX terlalu kompleks untuk tahap simulasi", "STRUCTURE_LIMIT"
                )
            if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise _quarantined(
                    "Ukuran isi XLSX setelah dibuka melebihi batas aman", "EXPANDED_SIZE_LIMIT"
                )
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise ExtractionRejected("Isi berkas tidak sesuai dengan format XLSX")
            if any(name.endswith("vbaproject.bin") for name in names):
                raise _quarantined(
                    "Dokumen yang mengandung makro tidak diizinkan", "MACRO_DETECTED"
                )
            if any(name.startswith("xl/externallinks/") for name in names):
                raise _quarantined(
                    "Tautan data luar pada Excel tidak diizinkan", "EXTERNAL_LINK_DETECTED"
                )
    except ExtractionRejected:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ExtractionRejected("XLSX tidak dapat dibaca dengan aman") from exc

    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        if len(workbook.sheetnames) > MAX_SPREADSHEET_SHEETS:
            raise _quarantined("Berkas Excel dibatasi maksimal 20 lembar", "SHEET_LIMIT")
        sheets: list[tuple[str, list[list[Any]]]] = []
        for worksheet in workbook.worksheets:
            if worksheet.max_column > MAX_SPREADSHEET_COLUMNS:
                raise _quarantined(
                    "Setiap lembar Excel dibatasi maksimal 100 kolom", "COLUMN_LIMIT"
                )
            rows = [
                [cell.value for cell in row]
                for row in worksheet.iter_rows(max_row=MAX_SPREADSHEET_ROWS + 1)
            ]
            sheets.append((worksheet.title, rows))
        workbook.close()
        return _spreadsheet_text(sheets)
    except ExtractionRejected:
        raise
    except Exception as exc:
        raise ExtractionRejected("XLSX tidak dapat dibaca dengan aman") from exc


def _extract_xls(payload: bytes) -> tuple[str, None, None, int]:
    if not payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise ExtractionRejected("Isi berkas tidak sesuai dengan format XLS")
    try:
        with OleFileIO(io.BytesIO(payload)) as container:
            stream_names = ["/".join(parts).lower() for parts in container.listdir()]
            if any(
                marker in name
                for name in stream_names
                for marker in ("vba", "_vba_project", "macros")
            ):
                raise _quarantined(
                    "Dokumen yang mengandung makro tidak diizinkan", "MACRO_DETECTED"
                )
            if sum(container.get_size(parts) for parts in container.listdir()) > 20 * 1024 * 1024:
                raise _quarantined(
                    "Ukuran isi XLS setelah dibuka melebihi batas aman", "EXPANDED_SIZE_LIMIT"
                )
        workbook = xlrd.open_workbook(file_contents=payload, on_demand=True)
        if workbook.nsheets > MAX_SPREADSHEET_SHEETS:
            raise _quarantined("Berkas Excel dibatasi maksimal 20 lembar", "SHEET_LIMIT")
        sheets: list[tuple[str, list[list[Any]]]] = []
        for sheet in workbook.sheets():
            if sheet.nrows > MAX_SPREADSHEET_ROWS:
                raise _quarantined(
                    "Setiap lembar Excel dibatasi maksimal 5.000 baris", "ROW_LIMIT"
                )
            if sheet.ncols > MAX_SPREADSHEET_COLUMNS:
                raise _quarantined(
                    "Setiap lembar Excel dibatasi maksimal 100 kolom", "COLUMN_LIMIT"
                )
            sheets.append((sheet.name, [sheet.row_values(index) for index in range(sheet.nrows)]))
        workbook.release_resources()
        return _spreadsheet_text(sheets)
    except ExtractionRejected:
        raise
    except Exception as exc:
        raise ExtractionRejected("XLS tidak dapat dibaca dengan aman") from exc


def extract_simulated_document(
    *, filename: str, content_type: str | None, payload: bytes
) -> ExtractedDocument:
    """Extract a bounded simulated file after extension, MIME, and structure checks."""
    safe_name = PurePath(filename.replace("\\", "/")).name
    extension = PurePath(safe_name).suffix.lower()
    if not _SIMULATED_NAME.match(safe_name):
        raise _quarantined("Nama berkas harus diawali 'SIMULASI'", "SIMULATION_LABEL_REQUIRED")
    if extension not in _ALLOWED_TYPES:
        raise _quarantined(
            "Gunakan berkas simulasi berformat TXT, MD, PDF, DOCX, XLSX, atau XLS",
            "UNSUPPORTED_FORMAT",
        )
    mime = (content_type or "").lower().split(";", 1)[0].strip()
    if mime not in _ALLOWED_TYPES[extension]:
        raise _quarantined("Jenis berkas tidak sesuai dengan nama berkas", "MEDIA_TYPE_MISMATCH")
    if not payload:
        raise ExtractionRejected("Berkas simulasi kosong")
    if len(payload) > MAX_FILE_BYTES:
        raise _quarantined("Ukuran berkas simulasi dibatasi maksimal 2 MB", "FILE_SIZE_LIMIT")

    sheets: int | None = None
    if extension in {".txt", ".md"}:
        text, pages, paragraphs = _extract_plain(payload)
    elif extension == ".pdf":
        text, pages, paragraphs = _extract_pdf(payload)
    elif extension == ".docx":
        text, pages, paragraphs = _extract_docx(payload)
    elif extension == ".xlsx":
        text, pages, paragraphs, sheets = _extract_xlsx(payload)
    else:
        text, pages, paragraphs, sheets = _extract_xls(payload)

    return ExtractedDocument(
        text=_validate_text(text),
        extension=extension,
        mime_type={
            ".txt": "text/plain",
            ".md": "text/markdown",
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ".xls": "application/vnd.ms-excel",
        }[extension],
        size_bytes=len(payload),
        checksum_sha256=sha256(payload).hexdigest(),
        page_count=pages,
        paragraph_count=paragraphs,
        sheet_count=sheets,
    )
